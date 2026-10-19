from pathlib import Path
import json, hashlib, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE

WORK=Path(__file__).parent; ROOT=WORK.parents[1]
old=Document(WORK/'v1-archiv/Projektübersicht.docx')
oldtables=[[[c.text for c in r.cells] for r in t.rows] for t in old.tables]
NAVY='#18384F';TEAL='#187F83';GRAY='#586775';PALE='#EFF4F6'
FONTS=Path('C:/Windows/Fonts')
def font(size,bold=False):return ImageFont.truetype(str(FONTS/('arialbd.ttf' if bold else 'arial.ttf')),size)
def diagram():
    im=Image.new('RGB',(1900,1060),'white');d=ImageDraw.Draw(im)
    labels=[('Weboberfläche und JSON API','Anmeldung, Erfassung und Bearbeitung'),('Fachschicht SDN.Leistung','Bedarf, Bescheid und Zahlungsworkflow'),('Transaktionssteuerung','TxManager, Tasks und Saga Journal'),('Oracle Datenbank','Atomarität, Sperren und Recovery')]
    for i,(title,body) in enumerate(labels):
        y=35+i*245
        d.rounded_rectangle((35,y,1310,y+165),radius=10,fill=PALE,outline='#CFDCE2',width=2)
        d.text((75,y+28),title,font=font(43,True),fill=NAVY)
        d.text((75,y+96),body,font=font(32),fill=GRAY)
        if i<3:
            d.line((672,y+180,672,y+225),fill=TEAL,width=6)
            d.polygon([(659,y+215),(685,y+215),(672,y+234)],fill=TEAL)
    d.rounded_rectangle((1430,35,1870,900),radius=10,fill='#E9F2F2',outline='#C6DADA',width=2)
    d.text((1480,100),'testsweet',font=font(44,True),fill=NAVY)
    for y,words in [(245,['Fachläufe und','Browserabläufe']),(460,['Automatische','Regressionen']),(675,['SQL Nachweise','und GLM Prüfung'])]:
        for k,line in enumerate(words):d.text((1480,y+k*48),line,font=font(31),fill=GRAY)
    for y in [115,360,850]:
        d.line((1430,y,1340,y),fill=TEAL,width=5);d.polygon([(1334,y),(1354,y-10),(1354,y+10)],fill=TEAL)
    im.save(WORK/'architektur-v2.png')
    im=Image.new('RGB',(1900,460),'white');d=ImageDraw.Draw(im)
    items=[('Automatische Tests','JUnit und Produktregression'),('Simulator','Sieben Pflichtkombinationen'),('GLM Prüfung','Ausgewählte vollständige Fälle'),('Zertifikat','Gebundene Belege und Export')]
    for i,(t,b) in enumerate(items):
        x=20+i*475
        d.ellipse((x+162,15,x+238,91),fill=TEAL)
        d.text((x+190,25),str(i+1),font=font(43,True),fill='white')
        d.text((x+20,140),t,font=font(34,True),fill=NAVY)
        for k,line in enumerate(textwrap.wrap(b,24)):d.text((x+20,200+43*k),line,font=font(29),fill=GRAY)
        if i<3:
            d.line((x+260,53,x+445,53),fill='#B6C8CE',width=4)
            d.polygon([(x+449,53),(x+432,44),(x+432,62)],fill='#B6C8CE')
    d.text((35,365),'Ein gemeinsamer Quellenstand und fortlaufende Hashkontrollen verbinden alle Prüfebenen.',font=font(31),fill=GRAY)
    im.save(WORK/'pruefkette-v2.png')
diagram()

doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(2);sec.right_margin=Cm(2);sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8)
sec.header_distance=Cm(.7);sec.footer_distance=Cm(.75);sec.different_first_page_header_footer=True
styles=doc.styles
for st in styles:
    for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Arial';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.line_spacing=1.13;styles['Normal'].paragraph_format.space_after=Pt(8)
styles['Title'].font.size=Pt(34);styles['Title'].font.bold=True;styles['Title'].paragraph_format.space_after=Pt(9)
styles['Subtitle'].font.size=Pt(17);styles['Subtitle'].font.italic=False;styles['Subtitle'].paragraph_format.space_after=Pt(9)
styles['Heading 1'].font.size=Pt(22);styles['Heading 1'].paragraph_format.space_after=Pt(15);styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Heading 2'].font.size=Pt(13);styles['Heading 2'].paragraph_format.space_before=Pt(12);styles['Heading 2'].paragraph_format.space_after=Pt(7)
styles['Heading 3'].font.size=Pt(12.5);styles['Heading 3'].font.bold=True;styles['Heading 3'].paragraph_format.space_before=Pt(12);styles['Heading 3'].paragraph_format.space_after=Pt(6)
styles['Caption'].font.size=Pt(8.5);styles['Caption'].paragraph_format.space_after=Pt(10)
for name in ['TOC 1','TOC 2','Front Title','Chapter Continuation']:
    if name not in styles:styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    styles[name].font.name='Arial';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Front Title'].font.size=Pt(22);styles['Front Title'].font.bold=True;styles['Front Title'].paragraph_format.space_after=Pt(15)
styles['Chapter Continuation'].base_style=styles['Heading 2'];styles['Chapter Continuation'].font.size=Pt(18);styles['Chapter Continuation'].font.bold=True;styles['Chapter Continuation'].paragraph_format.space_before=Pt(0);styles['Chapter Continuation'].paragraph_format.space_after=Pt(13)
styles['TOC 1'].font.size=Pt(11);styles['TOC 1'].font.bold=True;styles['TOC 1'].paragraph_format.space_before=Pt(10);styles['TOC 1'].paragraph_format.space_after=Pt(4)
styles['TOC 2'].font.size=Pt(10.5);styles['TOC 2'].paragraph_format.space_after=Pt(4);styles['TOC 2'].paragraph_format.left_indent=Cm(.55)
for name in ['Hyperlink','FollowedHyperlink']:
    if name in styles:styles[name].font.color.rgb=RGBColor(0,0,0);styles[name].font.underline=False
header=sec.header.paragraphs[0];header.add_run('SDN.Leistung  |  Projektdokumentation').font.size=Pt(8)
footer=sec.footer.paragraphs[0];footer.style=styles['Normal'];footer.paragraph_format.tab_stops.add_tab_stop(Cm(17),WD_ALIGN_PARAGRAPH.RIGHT)
r=footer.add_run('Projektübersicht  ·  Stand 15. Oktober 2026\t');r.font.size=Pt(8)
for instr in ['PAGE','NUMPAGES']:
    if instr=='NUMPAGES':footer.add_run(' / ').font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),instr)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'16');pr.append(sz);run.append(pr);txt=OxmlElement('w:t');txt.text='1';run.append(txt);fld.append(run);footer._p.append(fld)
doc.core_properties.title='Projektübersicht';doc.core_properties.subject='SDN.Leistung Entwicklungsverlauf und Qualitätssicherung';doc.core_properties.author='SDN.Leistung Projekt'
doc.core_properties.comments='Redaktionell und gestalterisch überarbeitete Projektdokumentation. Historischer Prüfstand 15.10.2026.'
def p(text,small=False,bold=False):
    pp=doc.add_paragraph();r=pp.add_run(text);r.bold=bold
    if small:r.font.size=Pt(9)
    return pp
def h(text,level=2):doc.add_heading(text,level)
def page(title,level=1):
    doc.add_page_break()
    if level==2:doc.add_paragraph(title,'Chapter Continuation')
    else:h(title,level)
def ref(text):p('Quellen  '+text,small=True)
def picture(name,width=17,alt=''):
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(7)
    sh=pp.add_run().add_picture(str(WORK/name),width=Cm(width));sh._inline.docPr.set('descr',alt)
def table(headers,rows,widths,font_size=9):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    for i,row in enumerate(t.rows):
        trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
        if i==0:trp.append(OxmlElement('w:tblHeader'))
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'18384F' if i==0 else ('F0F4F6' if i%2 else 'FFFFFF'));pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');mar.append(el)
            pr.append(mar);b=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');b.append(el)
            pr.append(b)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.05
                for r in pp.runs:r.font.size=Pt(font_size);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

# Titelblatt
pp=p('SDN.Leistung',bold=True);pp.paragraph_format.space_before=Pt(25);pp.paragraph_format.space_after=Pt(23)
doc.add_paragraph('Projektübersicht','Title')
doc.add_paragraph('Leistungsmanagement Grundsicherung\nund Transaktionssteuerung','Subtitle')
p('Entwicklung, Architektur und Qualitätssicherung',small=True)
picture('titelbild.png',17,'Abstraktes Architekturmotiv mit präzise geschichteten Elementen und einem durchgehenden Verbindungspfad.')
pp=p('Dokumentierter Projektverlauf\n24. September bis 15. Oktober 2026');pp.paragraph_format.space_before=Pt(20)
p('Projektstand 0.5.8  |  Abschlussnachweis vom 15. Oktober 2026',small=True)
p('Dokumentation der aktuellen Ausbaustufe',small=True)

page('Vorwort')
p('Diese Projektübersicht dokumentiert die Entwicklung von SDN.Leistung und seiner Transaktionssteuerung bis zum 15. Oktober 2026. Sie beschreibt den fachlichen und technischen Auftrag, ordnet die wesentlichen Entwicklungsschritte zeitlich ein und zeigt, wie sich der Funktionsumfang und die Qualitätssicherung gemeinsam entwickelt haben.')
p('Der früheste belegte Stand stammt vom 24. September 2026. Bereits diese Version enthielt einen Transaktionskern, eine Fachschicht für die Grundsicherung und eine Weboberfläche. In den folgenden drei Wochen wurden die Abläufe erweitert und systematisch gegen Fehler, konkurrierende Zugriffe und Unterbrechungen abgesichert. Den tatsächlichen Beginn der Vorarbeiten dokumentiert das Repository nicht.')
p('Eine besondere Rolle übernehmen die wiederholten Reviews und der Mitarbeiter-Simulator testsweet. Reviews hinterfragten Architektur, Berechnungen, Sicherheitsmechanismen und Prüfaussagen. testsweet ergänzte die Prüfung einzelner Komponenten um vollständige Arbeitsabläufe mit mehreren simulierten Sachbearbeitern und einer unmittelbaren Kontrolle der gespeicherten Daten.')
p('Die Dokumentation richtet sich an Leser, die den Projektverlauf und den erreichten Stand nachvollziehen möchten. Der Hauptteil erläutert die Zusammenhänge; die Anhänge enthalten das ausführliche Reviewregister und die Quellen. Grundlage sind die Git-Historie, die Projektdokumentation sowie die gespeicherten Review- und Laufbelege.')
h('Erreichter Stand')
p('Der vollständige Zertifikatslauf vom 15. Oktober endete um 22:40 Uhr erfolgreich. Er umfasste 395 Haupttests, 329 testsweet-Tests, 1.392 Simulatoraufträge und 200 ausgewählte GLM-Fälle. Im vorgesehenen Prüfumfang wurden keine Fehler oder Befunde gemeldet. Die Übersicht erläutert, welche Prüfungen dieses Ergebnis tragen und welche Modellgrenzen weiterhin bestehen.')
p('SDN.Leistung ist ein Technologie- und KI-Showcase. Diese Dokumentation beschreibt die abgeschlossene aktuelle Ausbaustufe und ihren Prüfstand; sie erteilt keine behördliche Produktionsfreigabe.')
ref('Q01–Q04 und Q24–Q26. Alle Uhrzeiten beziehen sich auf Europe/Berlin.')

doc.add_page_break();doc.add_paragraph('Inhaltsverzeichnis','Front Title')
doc.add_paragraph('[[TOC]]')

page('Projektauftrag und Zielsetzung')
p('SDN.Leistung verbindet eine eigene Transaktionssteuerung mit einer Fachanwendung für die Grundsicherung im Alter und bei Erwerbsminderung. Ausgangspunkt war die Frage, wie sich ausgewählte Aufgaben eines klassischen Transaktionsmonitors mit Java und Oracle umsetzen lassen. Im Mittelpunkt stehen die sichere Ausführung kurzer Arbeitseinheiten, kontrollierte Wiederholungen bei Sperrkonflikten, die Begrenzung paralleler Aufgaben und der Wiederanlauf längerer Geschäftsvorgänge.')
p('Die Fachanwendung macht diese technischen Anforderungen an einem zusammenhängenden Prozess überprüfbar. Eine Bedarfsgemeinschaft wird erfasst, ihre SozialFälle werden berechnet, ein Bescheid wird erteilt und eine Zahlung angeordnet. Jede Stufe verändert Daten und erzeugt Anforderungen an Konsistenz, Berechtigung und Nachvollziehbarkeit. Ein Fehler kann einen Betrag, einen Status oder den gesamten weiteren Ablauf betreffen.')
h('Aufgabenteilung zwischen Datenbank und Anwendung')
p('Das Architekturprinzip weist Oracle und der Java-Anwendung klar getrennte Aufgaben zu. Oracle liefert die grundlegenden Transaktionseigenschaften, die Sperrverwaltung, das Redo-Logging und die Datenbank-Recovery. Die Anwendung steuert den Beginn und Abschluss von Transaktionen, entscheidet über zulässige Wiederholungen und koordiniert mehrstufige Geschäftsvorgänge. Die Eigenentwicklung konzentriert sich damit auf die Steuerung über dem vorhandenen Datenbankkern.')
h('Fachlicher und technischer Umfang')
p('Die Fachschicht verarbeitet Regelbedarfe, Mehrbedarfe, Unterkunftskosten, Einkommen und Vermögen im dokumentierten Produktmodell. Sie ergänzt die Berechnung um Bescheide, Zahlungsanordnungen, Freigaben, Kassenübergabe und Storno. Rollen und Sitzungen regeln den Eintritt in die Anwendung; Audit und Einsichtsprotokolle machen ausgewählte Vorgänge nachvollziehbar.')
p('Technisch verwendet das Projekt Java mit Zielversion 21, Maven und Oracle. Die Webschicht basiert auf dem HTTP-Server des JDK. JSON-Verarbeitung, Transaktionssteuerung, Saga, Konfiguration und Monitoring sind eigene Bausteine. Für die Verbindungsverwaltung stehen ein minimaler Pool und Oracle UCP zur Verfügung.')
h('Maßstab für den Projektabschluss')
p('Der Abschluss der aktuellen Ausbaustufe umfasst einen implementierten und lauffähigen Demonstrationsablauf sowie einen vollständigen, an den Quellenstand gebundenen Prüfabschluss. Ein vollständiger Ersatz aller Funktionen eines Großrechner-Transaktionsmonitors gehört ebenso wenig zum dokumentierten Umfang wie ein behördlicher Produktiveinsatz. Die konkreten fachlichen Grenzen werden im Kapitel Ergebnis und weiterer Entwicklungsbedarf zusammengeführt.')
ref('Q02–Q04; ergänzend doc/architektur.md und doc/leistungsmanagement.md.')

page('Architektur und Funktionsumfang')
p('Die Architektur verbindet Oberfläche, Fachlogik und Transaktionssteuerung mit einem gemeinsamen Oracle-Datenbestand. testsweet prüft diese Schichten über unterschiedliche Zugangswege und kontrolliert ihre persistierten Wirkungen.')
picture('architektur-v2.png',17,'Vier Schichten: Weboberfläche, Fachschicht, Transaktionssteuerung und Oracle; testsweet prüft Fach- und Webwege sowie Datenbankzustände.')
doc.add_paragraph('Abbildung 1  Zusammenspiel der Anwendungsschichten und des Prüfwerkzeugs','Caption')
h('Kurze Transaktionen und lange Geschäftsvorgänge')
p('Zusammengehörende kurze Änderungen werden durch den Transaktionsmanager gemeinsam bestätigt oder bei Fehlern zurückgerollt. Nur definierte Oracle-Sperrkonflikte erlauben einen erneuten Versuch. Für längere Abläufe speichert eine Saga den Fortschritt mehrerer kurzer Transaktionen im Journal. Wiederholbare Kompensationen behandeln bereits ausgeführte Schritte nach einem Fehler oder Absturz.')
h('Schutz der fachlichen Verarbeitung')
p('Zeilensperren, Versionsprüfungen und Idempotenzmerkmale schützen gegen Konkurrenz und Doppelverarbeitung. Zahlungsanordnungen durchlaufen einen Vier-Augen-Workflow. Bescheide speichern ihren Berechnungs- und Dokumentstand in versionierten Snapshots. Bei fehlender Commitquittung meldet die Anwendung einen unbekannten Ausgang, damit eine möglicherweise wirksame Änderung vor einer Wiederholung geklärt wird.')
ref('Q03, Q17–Q21; doc/architektur.md. Die Grafik zeigt die logische Gliederung, keinen vollständigen Aufrufgraphen.')

page('Projektverlauf im Überblick')
p('Zwischen dem ersten belegten Stand und dem Abschlussnachweis liegen 21 Kalendertage. Die Entwicklung lässt sich in drei Phasen gliedern: der Ausbau des Fachsystems, die Prüfung vollständiger Arbeitsabläufe und die Absicherung der gesamten Nachweiskette.')
picture('timeline.png',17,'Chronologische Übersicht vom 24. September bis zum 15. Oktober 2026 mit sechs Entwicklungsschritten.')
doc.add_paragraph('Abbildung 2  Wesentliche Meilensteine des dokumentierten Projektverlaufs','Caption')
p('Die Zeitleiste ordnet Ereignisse chronologisch; ihre Abstände bilden keine Entwicklungsdauer ab. Detaillierte Zeitangaben in der folgenden Darstellung sind Git-Integrationszeiten. Reviewtermine und Prüflaufzeiten werden gesondert ausgewiesen.',small=True)
ref('Q01, Q03 und Q05–Q26.')

page('Vom ersten Stand zum erweiterten Fachsystem',2)
p('Die belegte Entwicklung setzt am 24. September 2026 mit Version 0.2.0 ein. Dieser erste gespeicherte Stand enthielt bereits Transaktionsmanager, Wiederholungsregeln, Tasksteuerung und Saga-Journal sowie Fachberechnung, Zahlungsworkflow und Weboberfläche. Die damalige Roadmap verzeichnete einen erfolgreichen Build mit 23 Tests und Prüfungen gegen eine echte Oracle-Datenbank. Noch am selben Abend folgte Version 0.3.0 mit Anmeldung, Rollen, Neuanlage, Kassenlauf und Storno.')
p('Am 10. Oktober wurde das System in mehreren Versionen ausgebaut. Version 0.4.0 ergänzte SEPA-Bankverbindungen, das Vier-Augen-Prinzip, das Bescheidwesen und eine Antrags-Saga. Zugleich kamen Betriebsbausteine wie Oracle UCP, Monitoring, externalisierte Konfiguration und Audit hinzu. Damit entstand ein durchgängiger fachlicher Ablauf mit zugehörigen Verwaltungs- und Überwachungsfunktionen.')
table(['Integration','Version','Wesentliche Ergänzung'],[
('24.09. 21:24','0.2.0','Erste belegte funktionsfähige Umsetzung'),('24.09. 21:43','0.3.0','Anmeldung, Neuanlage, Kassenlauf und Storno'),('10.10. 10:08','0.4.0','SEPA, Bescheid, Antrags-Saga und Betriebsbausteine'),('10.10. 12:47','0.5.0','Erweiterte Fachregeln sowie Pool-, Task- und Saga-Härtung'),('10.10. 13:15','0.5.1','Audit-Funde zu XSS, Eingabegrenzen und Webschutz'),('10.10. 14:11','0.5.2','Fachkorrekturen, Monatsprüfung und Saga-Kompensation'),('10.10. 17:49','0.5.3','Bescheidtext und eindeutiger Monatsindex'),('10.10. 19:41','0.5.4','Rechtsquellen und eingefrorener Berechnungssnapshot')],[3.0,1.6,12.4])
p('Die Versionen 0.5.0 bis 0.5.4 standen zugleich im Zeichen früher Review-Nacharbeiten. Berechnungsregeln wurden erweitert oder korrigiert; Transaktionsabschluss, Verbindungspools, Taskabbruch und Saga-Recovery wurden genauer abgesichert. Ein Rechtsquellenverzeichnis und versionierte Berechnungssnapshots verbesserten die Nachvollziehbarkeit fachlicher Entscheidungen.')
p('Bereits am 10. Oktober war der fachliche Ablauf vom Antrag bis zur Auszahlung dokumentiert umgesetzt. Die folgenden Arbeiten vertieften vor allem die Absicherung seines Verhaltens in Grenz- und Fehlerfällen. Für die frühen Audit- und Gegenprüfungen sind Korrekturen in der Versionsgeschichte belegt; eine vollständige Einzelbilanz mit Prüferzuordnung liegt nicht vor.',small=True)
ref('Q01–Q04; Commits 41f23eb bis cc6da3d.')

page('Von Arbeitsabläufen zum vollständigen Prüfnachweis',2)
h('11 und 12 Oktober 2026')
p('Am 11. Oktober ergänzte eine Szenario-Regression die bisherigen Prüfungen. Zehn Nutzerabläufe verbanden mehrere Bearbeitungsschritte und Sachbearbeiter miteinander. Bewertet wurde der persistierte Datenbankzustand: Beträge, Statuswechsel, Zuordnungen und Journalwirkungen mussten den vorgegebenen Erwartungen entsprechen. Damit rückte die Zusammenarbeit der Komponenten stärker in den Mittelpunkt.')
p('Am selben Tag entstand testsweet als eigenständiger Mitarbeiter-Simulator. Im ersten Modus bearbeiteten simulierte Sachbearbeiter einen gemeinsamen Auftragstopf über die Fachschicht. Protokolle sowie PDF- und Markdown-Berichte machten die Abläufe und Datenbankprüfungen nachvollziehbar. Der anschließend ergänzte Selenium-Modus bezog die Weboberfläche ein. Positiv-, Negativ-, Last- und Stresstests erweiterten die Prüfung auf erwartete Ablehnungen, größere Auftragsmengen und Änderungen derselben Fälle in mehreren Wellen.')
p('Die Reviews vom 12. Oktober führten zu weiteren Änderungen im Produkt und im Prüfwerkzeug. Neben Sicherheits- und Robustheitskorrekturen kamen eine datenbankgestützte Satztabelle, die Vermögensfortschreibung des Demomodells und ein Einsichtsprotokoll hinzu. Die Versionen 0.5.5 bis 0.5.7 wurden gemeinsam integriert. Version 0.5.8 ergänzte monatliche Bewilligungspläne und weitere Regressionen. Auch Browserprofile, Wiederaufnahme und globale SQL-Prüffehler wurden untersucht.')
h('13 bis 15 Oktober 2026')
p('Ab dem 13. Oktober erhöhte sich die Prüftiefe von testsweet. Generierte Konstellationen ergänzten kuratierte Szenarien; feste fachliche Grenzfälle erweiterten die gezielte Abdeckung. Eine echte GLM-5.3-Nachprüfung bewertete ausgewählte Fälle anhand vollständiger benannter Dokumenttexte. Gespeicherte Anfragen, Antworten und Quellsnapshots hielten fest, welche Unterlagen dem Modell tatsächlich übergeben worden waren.')
p('Das Zertifikatswerkzeug verband automatisierte Tests, die Simulator-Matrix, die Modellstichprobe und den Dokumentexport. Quellen- und Buildbindung verhinderten, dass ältere erfolgreiche Ergebnisse als aktueller Nachweis für geänderten Code erscheinen konnten. Der erste vollständige erfolgreiche Gesamtlauf fand am 14. Oktober statt.')
p('Weitere Gesamtprüfungen am 15. Oktober führten zu Produktregressionen und strengeren Nachweisen. Die Prüf-Runtime wurde laufgebunden kopiert, Testdaten erhielten eigene UUID-Namensräume und die Auswertung von Modellantworten wurde vereinheitlicht. Der abschließende Gesamtlauf von 21:29 bis 22:40 Uhr bestätigte schließlich den vorgesehenen vollständigen Zertifikatsumfang.')
ref('Q01, Q03, Q15–Q28. Die Produktversion blieb bei den späteren Nachbesserungen 0.5.8.')

page('Reviews und Fehlerbehebung')
p('Die Reviewstrategie entwickelte sich mit dem Projekt. An die Prüfung von Architektur und Fachrechnung schlossen sich Gegenproben zu Berechtigungen, Konkurrenz, Recovery und Nachweisqualität an. Insgesamt sind zehn eigenständige Primärberichte vom 12., 13. und 15. Oktober dokumentiert. Teilberichte, Fixberichte und Testwiederholungen werden im Register gesondert eingeordnet.')
h('Erste Reviewwelle am 12 Oktober')
p('Der erste ausführliche Gesamtreview zeigte trotz grüner Ausgangsläufe 75 nummerierte Feststellungen. Davon waren 60 problemorientierte Punkte und 15 Informationspunkte, die teilweise bereits korrekte Eigenschaften bestätigten. Besonders schwer wogen Sozialdaten im Audit-Log, Passwörter in Testberichten und unterschiedliche Berechnungen im manuellen Fachweg und in der Antrags-Saga.')
p('Alle drei Blocker und neun kritischen Funde wurden anschließend als behoben dokumentiert. Die Hauptsuite wuchs beim ersten Korrekturstand von 136 auf 195 Tests. Eine zusätzliche Nachprüfung erfasste elf weitere Punkte. Deren Korrekturen stärkten unter anderem die monatsgenaue Bewilligung, die Absicherung gegen gleichzeitige Historisierung und die Bindung wiederholter Zahlungsanforderungen.')
h('Prüfung des Prüfwerkzeugs am 13 Oktober')
p('Vier weitere Primärreviews untersuchten, ob ein fehlerhafter oder unvollständiger Nachweis trotzdem als bestanden erscheinen konnte. Ihre Listen enthalten zusammen 66 nummerierte Hauptbefundnennungen und zusätzliche Randbeobachtungen. Die Berichte überschneiden sich. Ein wesentlicher Kritikpunkt war der Zirkelschluss, erwartete Beträge mit derselben Produktionsrechnung zu erzeugen, deren Korrektheit geprüft werden sollte.')
p('Separate Rechenreferenzen, vorgegebene Handwerte und strengere Zustandsvergleiche machten die Prüfkette unabhängiger. Echte GLM-Aufrufe ersetzten die irreführende Bezeichnung einer lokalen Nachrechnung als Modellprüfung. Erkannte Befunde, fehlende Pflichtprüfungen und technische Abbrüche flossen gemeinsam in das Gesamturteil ein.')
h('Weitere Gesamtprüfungen am 15 Oktober')
p('Fünf weitere Gesamtberichte führten zu 41 priorisierten Befundnennungen und neun zusätzlichen Info-Beobachtungen. Grok, Antigravity, DeepSeek und GLM-5.3 sind bei einzelnen Berichten als Herkunft belegt. Die Gegenprüfungen bestätigten neue Schwachstellen und präzisierten oder widerlegten zugleich einzelne Reparaturforderungen. Das ausführliche Register im Anhang weist Datum, Zuordnung und Befundzahl aus.')
ref('Q05–Q21. Die Befundnennungen verschiedener Wellen sind keine entdoppelte Gesamtfehlerzahl.')

page('Korrekturen mit dauerhaftem Nachweis',2)
p('Die Verbesserung der Fehlersicherheit zeigt sich an konkreten Fehlerpfaden und den daraus entstandenen Regressionen. Die folgenden Beispiele verdeutlichen, wie ein Reviewfund in eine überprüfbare Korrektur überführt wurde.')
h('Berechtigung beim Zahlungsstorno')
p('Der Gesamtreview vom 15. Oktober reproduzierte einen Storno über alternative ID-Schreibweisen, obwohl der Aufrufer nur Sachbearbeiterrechte hatte. Eine zusätzliche Rollenprüfung unmittelbar vor dem Serviceaufruf schloss den Pfad. Die Regressionen prüfen die vollständige Kette über HTTP, Fachservice und Oracle, einschließlich kodierter und nichtkanonischer IDs.')
h('Fachliche Grenzfälle und eingefrorene Bescheide')
p('Bei bestimmten Haushalten mit vier oder mehr Kindern übernahm die Berechnung eine ungünstigere Alternative. Auch ein konfigurierbarer zentraler Warmwasserabzug konnte den Bedarf fehlerhaft mindern. Mehrere Referenzrechner teilten diese Annahmen. Die Korrektur wurde deshalb durch vorab handgerechnete Grenzfälle ergänzt. Neue Bescheidsnapshots speichern außerdem das vollständige Erteilungsdokument, sodass spätere Stammdatenänderungen den historischen Nachdruck nicht verändern.')
h('Unbekannter Ausgang statt falscher Rollbackzusage')
p('Ein Taskabbruch oder eine fehlende Commitquittung beweist nicht, dass eine innere Fachtransaktion zurückgerollt wurde. Die Gegenprüfungen führten zu einem ausdrücklich unbekannten Ausgang und einer HTTP-409-Antwort. Eine möglicherweise wirksame Aktion wird nicht blind wiederholt. Neue Regressionen behandeln auch einen tatsächlich ausgeführten Commit mit synthetisch verlorener Quittung und den fatalen Treiberfehlerpfad.')
h('Wiederanlauf und historische Journale')
p('Saga-Kompensationen werden über gespeicherte Schrittnamen gebunden. So führt ein späterer Definitionswechsel nicht zur falschen Zuordnung über einen aktuellen Listenindex. Historische namenlose PENDING-Einträge benötigen besondere Vorsicht: Sie beweisen nicht sicher, dass keine Fachwirkung eingetreten ist. Ungeklärte Fälle bleiben mit erhaltenen Daten bestehen und können über ein kontrolliertes Offline-Verfahren untersucht werden.')
h('Nachweise unmittelbar an der Fachwirkung')
p('Die Storno-Nachprüfung liest Vermögen und Restbudget heute unmittelbar aus der Datenbank, bevor eine weitere Änderung erfolgt. Ein Browserproblem mit ersetzten Listenzeilen wird nur beim lesenden Öffnen erneut versucht; bestätigte Schreibschritte werden nicht automatisch wiederholt. Eigene Laufkennungen begrenzen Cleanup auf eindeutig zugeordnete Testdaten.')
p('Die Gegenprüfung des Grok-Berichts fasst diese differenzierte Bearbeitung beispielhaft zusammen: 16 Ausgangspunkte führten zu zwölf Korrekturen, drei nicht bestätigten Reparaturforderungen und einer offenen Zuständigkeitsgrenze. Die fachliche Prüfung der Forderung bleibt damit Teil der Fehlerbehebung.')
ref('Q10, Q17–Q21. Die neuen Tests sichern konkrete Eigenschaften; ihre Anzahl misst keine allgemeine Fehlerfreiheit.')

page('testsweet als eigenständiges Prüfwerkzeug')
p('testsweet entstand aus der Anforderung, vollständige Arbeitsabläufe mehrerer Sachbearbeiter nachvollziehbar zu prüfen. Das Werkzeug besitzt ein eigenes Maven-Projekt und eine eigene Testsuite. Es verteilt Aufträge auf simulierte Mitarbeiter, protokolliert die Aktionen und kontrolliert anschließend die gespeicherten Wirkungen in Oracle.')
p('Die erste Einführung ist am 11. Oktober um 11:13 Uhr belegt. Noch am selben Tag kam ein zweiter Zugangsweg hinzu: Selenium steuerte die echte Weboberfläche in parallelen Chrome-Fenstern. Damit lässt sich derselbe fachliche Prozess unmittelbar über die Fachschicht und aus Sicht eines angemeldeten Anwenders untersuchen.')
h('Sieben verbindliche Prüfkombinationen')
table(['Zugang','Kategorie','Prüfzweck'],[
('Fachschicht','Positiv','Erfolgreicher Ablauf mit vorgegebenem persistiertem Endzustand'),('Fachschicht','Negativ','Erwarteter Fachfehler und unveränderter Zustand am Abwehrschritt'),('Fachschicht','Last','Viele schlanke Aufträge mit begrenzter funktionaler Parallelprüfung'),('Fachschicht','Stress','Zwei Wellen über dieselben Fälle mit Phasenbarriere'),('Weboberfläche','Positiv','Fachlicher Ablauf über Browser und Oberfläche'),('Weboberfläche','Negativ','Erwartete Abwehr über die Oberfläche mit Zustandskontrolle'),('Weboberfläche','Last','Begrenzte parallele Browserabläufe')],[3.2,2.5,11.3])
h('Breitere Konstellationen und nachvollziehbare Berichte')
p('Ein Kombinationsgenerator ergänzt kuratierte Fälle um unterschiedliche Haushalte, Merkmale, Einkommen, Vermögen, Bearbeitungswege und Verstoßkombinationen. Eine Signaturmenge verhindert Wiederholungen. Feste, unabhängig vorgegebene Fachgrenzfälle sichern besonders relevante Konstellationen ab und werden bei der GLM-Auswahl bevorzugt.')
p('Die Laufberichte verbinden Bearbeitungsschritte, SQL-Abfragen, Sollwerte und tatsächliche Ergebnisse. PDF dient der lesbaren Dokumentation; Markdown ermöglicht die maschinelle Nachprüfung. Ein Negativfall darf erfolgreiche Vorbereitungsschritte enthalten. Entscheidend ist, dass die abgewehrte Aktion den betroffenen Zustand nicht unzulässig verändert.')
ref('Q01, Q06, Q22 und Q24. Last und Stress sind funktionale Prüfungen im ausgeführten Umfang, kein Kapazitätsnachweis.')

page('Unabhängige Prüfungen und belastbare Belege',2)
p('Die Reviews vom 13. Oktober machten deutlich, dass auch das Prüfwerkzeug seine eigenen Voraussetzungen nachweisen muss. Der erste testsweet-Review erfasste 20 bestätigte Befunde, die laut Nachtrag korrigiert und gegengeprüft wurden. Die weiteren Arbeiten verbanden unabhängige Sollwerte mit einer strengeren Auswertung aller Prüfschichten.')
picture('pruefkette-v2.png',17,'Vier aufeinander aufbauende Prüfebenen: automatische Tests, sieben Simulatorläufe, GLM-Prüfung und Zertifikat. Alle sind an denselben Quellenstand gebunden.')
doc.add_paragraph('Abbildung 3  Vom automatischen Test zum gebundenen Gesamtnachweis','Caption')
h('Rechnung und fachliches Urteil')
p('Separate BigDecimal- und BigInt-Referenzen rechnen die dokumentierten Zahlen nach. Handgerechnete Beispiele und gezielte Gegenproben prüfen zusätzlich die zugrunde liegenden Annahmen. Erwartete Fehler müssen tatsächlich auftreten; falsche Fehlergründe, zusätzliche Ergebniszeilen und SQL-Ausnahmen verhindern den Erfolg.')
p('GLM-5.3 erhält Prüfauftrag, bereitgestellte Gesetzesgrundlagen und vollständige ausgewählte Fallgruppen als benannte Dokumenttexte. Seine Antwort muss Fälle, Rechenblöcke und Prüfpunkte vollständig abdecken. Transportnachweis, formale Vollständigkeit und fachliches Urteil bleiben getrennt sichtbar. Das Modell ergänzt die automatischen Prüfer um eine zusätzliche Fallprüfung.')
h('Quellenbindung und isolierte Ausführung')
p('JUnit, Simulator, GLM und Zertifikat müssen zum selben erfassten Quellenstand gehören. SHA-256-Inventare, gespeicherte Anfragen und Antworten sowie eine gehasht kopierte Prüf-Runtime sichern die Zuordnung. Fortlaufende Kontrollen erkennen Änderungen auch während längerer Modellprüfungen. UUID-Namensräume halten Bereinigung und globale SQL-Prüfungen auf die eigenen Laufdaten begrenzt.')
p('Fehlerhafte, abgebrochene oder veraltete Pflichtnachweise verhindern ein positives Gesamturteil. Frühere Fehlversuche bleiben erhalten: etwa fehlende SEPA-Rohwerte, ein veralteter Browserknoten, eine ungeeignete Saga-Testfixture und ein während der Prüfung ersetztes Buildartefakt. Die Korrekturen wurden durch neue Läufe nachgewiesen; alte Ergebnisse wurden nicht nachträglich umgedeutet.')
ref('Q06–Q09, Q17, Q19–Q25 und Q27–Q28.')

page('Abschlussstand am 15 Oktober 2026')
p('Der jüngste vollständige Zertifikatsgesamtlauf startete um 21:29:01 Uhr und endete um 22:40:32 Uhr erfolgreich. Alle 13 Phasen wurden abgeschlossen. Der veröffentlichte Zustand lautet READY mit Urteil BESTANDEN und null Findings. Die folgende Matrix beschreibt den tatsächlichen Umfang dieses Nachweises.')
table(*[oldtables[9][0],oldtables[9][1:],[5.0,3.3,8.7]])
p('Die 200 GLM-Fälle stammen aus erfolgreichen Simulatorberichten und bilden eine ausgewählte zusätzliche Prüfung. Sie werden nicht als weitere unabhängige Simulatoraufträge gezählt. Ebenso werden die Zahlen von Java-, Node- und Python-Prüfungen nicht zu einer gemeinsamen Qualitätsquote addiert.',small=True)
h('Formaler Abschluss der Modellprüfung')
p('GLM prüfte genau 100 positive und 100 negative Fälle in 67 Gruppen. 59 Gruppen waren beim ersten Versuch gültig, sieben nach einer Wiederholung und eine nach zwei Wiederholungen. Insgesamt erfolgten neun erneute Sendungen mit protokollierten Korrekturhinweisen; keine Gruppe scheiterte endgültig. Die ursprünglichen Antworten blieben erhalten.')
h('Nachvollziehbarkeit des Quellenstands')
p('Der Nachweis gehört zum Arbeitsbaum auf Basiscommit 37abfb6 einschließlich gehashter lokaler Änderungen. Eine erneute Kontrolle der 288 inventarisierten Quellen und der fünf veröffentlichten Dateien ergab bei Erstellung der Übersicht keine Hashabweichung. Das Ergebnis lautet daher präzise: null beobachtete Fehler und Befunde im dokumentierten Prüfumfang.')
ref('Q24–Q26. Gesamtlauf 3a0bbe85-bd68-4739-8674-fd92b3329223; Zertifikat b805c5b6-20b0-4215-83a5-26900a2f31e4.')

page('Ergebnis und weiterer Entwicklungsbedarf')
p('Die aktuelle Ausbaustufe verbindet einen lauffähigen fachlichen Ablauf mit einer mehrstufigen, dokumentierten Qualitätssicherung. Die Entwicklung reicht vom Transaktionskern über vollständige Mitarbeiterabläufe bis zur Prüfung der Nachweiskette selbst. Der erfolgreiche Gesamtlauf vom 15. Oktober markiert den Abschluss dieses vorgesehenen Umfangs.')
h('Woran sich die Verbesserung erkennen lässt')
p('Bestätigte Fehlerpfade wurden korrigiert und durch gezielte Regressionen erneut erreichbar gemacht. Dazu gehören unberechtigte Stornos, gemeinsame falsche Rechenannahmen, unklare Commit-Ausgänge und problematische Wiederanlaufzustände. Breitere Szenarien, direkte Datenbankkontrollen und unabhängig vorgegebene Sollwerte ergänzen diese Einzelprüfungen.')
p('Auch die Erfolgsaussage selbst wurde strenger. Ein vollständiges Zertifikat verlangt heute passende Quellen- und Buildstände, erfolgreiche Pflichtprüfungen und unveränderte Belege. Ein positiver Teilbefund kann einen Fehler oder eine fehlende Pflichtprüfung nicht überdecken. Die Verbesserung ist damit an konkreten geschlossenen Fehlerpfaden und strengeren Nachweisen erkennbar; eine statistische Restfehlerquote wurde nicht bestimmt.')
h('Bekannte Grenzen')
p('Objekt- und Teamzuständigkeiten sind nicht implementiert. Die gemeinsame Bearbeitung des Demobestands entspricht dem aktuellen Modell; ein Betrieb mit getrennten Zuständigkeiten benötigt eine Erweiterung. Auch eine beleggebundene Vermögenshistorie und ein strenger Modus mit durchgehender Bescheidpflicht bleiben gesonderte Fachaufgaben. Ein lediglich rechnerischer Vermögensverbrauch ist kein allgemeiner rechtlicher Reparaturmechanismus.')
p('Der aktuelle Abschlusslauf wurde mit Temurin 25.0.4.1, Compilerziel Java 21 und Oracle Free 23.26.3 erzeugt. Er belegt keinen erneuten Gesamtlauf auf einer Java-21-Laufzeit oder Oracle 19c Enterprise Edition. Last- und Stressfälle untersuchen definierte Parallelabläufe; sie liefern keinen Dauerlast-, Sättigungs- oder Kapazitätsnachweis. GLM ist eine zusätzliche Stichprobenprüfung und keine amtliche fachliche Freigabe.')
h('Einordnung des Abschlusses')
p('Die Dokumentation weist zehn ausführliche Primärreviews, weitere Gegenprüfungen und einen vollständigen erfolgreichen Abschlusslauf nach. Wegen überlappender Befunde, Dokumentationspunkten und teilweise widerlegter Reparaturforderungen wird keine fiktive Gesamtzahl verschiedener Fehler angegeben. Der belegte Fortschritt liegt in der nachvollziehbaren Bearbeitung der Funde und in der zunehmenden Absicherung des definierten Produktmodells.')
ref('Q04, Q10, Q17–Q26.')

page('Reviewregister')
p('Das Register enthält zehn eigenständige Primärberichte. Die Reihenfolge innerhalb eines Tages dient der Zuordnung und ist keine vollständig belegte minutengenaue Ablaufreihenfolge. Die KI-Herkunft wird nur bei ausdrücklicher Dokumentation angegeben.')
h('12 und 13 Oktober 2026')
table(oldtables[2][0],oldtables[2][1:],[2.3,7.0,7.7])
p('Bis zum Ende des 13. Oktober liegen damit fünf Primärreviews vor. Die erste Welle enthält 60 problemorientierte Punkte und 15 Info-Feststellungen. Die vier Folgeberichte nennen 66 nummerierte Hauptpunkte; der Abendbericht enthält zusätzlich 30 niedrige Beobachtungen. Überschneidungen verhindern eine Addition zu einer Zahl eindeutiger Fehler.')
h('Zählkorrekturen und Nachbesserungen')
p('Die ID-Liste CR037–060 enthält 24 Minor-Punkte, obwohl die Summenzeile des ersten Berichts 23 nennt. Das Register folgt den IDs und damit 75 Gesamtfeststellungen. Die statische Fixverifikation vom 13. Oktober führt tatsächlich CR201–210 aus; die abweichende Einleitung bis CR213 wird nicht als zusätzliche Befundliste gezählt.')
p('Zusätzlich dokumentiert sind die Behebung aller drei Blocker und neun kritischen Funde des ersten Reviews, elf neue Nachprüfungsbefunde N01–N11 sowie die gemeinsame Bearbeitung der teilweise überlappenden CR101–120 und R01–R16.')
ref('Q05–Q09, Q15–Q16 und Q27.')

page('Gesamtprüfungen und Gegenprüfungen vom 15 Oktober',2)
table(oldtables[3][0],oldtables[3][1:],[3.8,4.4,8.8])
p('Mit diesen fünf Berichten wächst das Register auf zehn Primärreviews. Die Tagesbilanz umfasst 41 priorisierte Befundnennungen und neun Info-Beobachtungen. Wiederbefunde und bekannte Modellgrenzen sind enthalten; auch diese Bilanz ist keine Zahl eindeutiger neuer Produktfehler.')
h('Dokumentierte Bearbeitungsbilanz')
p('Die zehn GR-Hauptfunde wurden im zugehörigen Fixbericht einzeln korrigiert und mit Regressionen verknüpft. Beim Grok-Bericht führten 16 Ausgangspunkte zu zwölf Korrekturen, drei nicht bestätigten Forderungen und einer offenen Zuständigkeitsgrenze. Die Antigravity-Gegenprüfung korrigierte Testisolation und Nachweislücken und wies eine vorgeschlagene Produktreparatur zurück.')
p('Die GLM-Gegenprüfung bearbeitete alle fünf priorisierten Punkte, darunter eine Testlücke und einen latenten Fehler. Sie entdeckte zusätzlich das Integritätsrisiko historischer namenloser PENDING-Journale. Die DeepSeek-Gegenprüfung korrigierte vier Hauptpunkte ganz oder teilweise: Commitquittung, Lastaussage, Cleanup und Eingabevertrag. Andere Forderungen wurden als bekannte Grenzen, nicht bestätigte Sollzustände oder bereits korrigierte Punkte eingeordnet.')
ref('Q10–Q14 und Q17–Q21. Antigravity bezeichnet das Werkzeug; das zugrunde liegende Modell ist nicht belegt.')

page('Quellen und Nachweise')
p('Die Quellenkennungen verknüpfen die Darstellung mit der Projektdokumentation. Alle Pfade beziehen sich auf das Repository LeistungsmanagementGrundsicherung. Die Quellen dienen der historischen Einordnung; für die Gültigkeit des aktuellen Zertifikats bleiben dessen maschinelle Belege maßgeblich.')
table(oldtables[11][0],oldtables[11][1:],[1.4,15.6],8.7)
p('Historische Testzahlen beschreiben ihren jeweiligen Quellenstand. Reviewzahlen werden aus konkreten Befundlisten übernommen; Info-Punkte und wiederholte Meldungen werden nicht pauschal als zusätzliche Softwarefehler gezählt.',small=True)

page('Weitere Nachweise und Bildnachweis',2)
table(oldtables[12][0],oldtables[12][1:],[1.4,15.6],8.7)
p('Titelbild: mit dem integrierten Bildgenerierungswerkzeug erstellte redaktionelle Illustration. Motivvorgabe: präzise geschichtete dunkelblaue und gläserne Architekturelemente, verbunden durch einen türkisfarbenen Pfad vor hellem Hintergrund. Das Motiv symbolisiert die gegliederte Verarbeitung; es ist kein Systemdiagramm. Bilddatei und vollständiger Generierungsprompt liegen unter output/projektuebersicht/titelbild.png und titelbild-nachweis.json.',small=True)

output=ROOT/'doc/Projektübersicht.docx';doc.save(output)
print(output)
