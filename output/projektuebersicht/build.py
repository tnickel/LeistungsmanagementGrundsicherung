from pathlib import Path
import json, hashlib, textwrap, subprocess
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).parent
DOC = ROOT / 'doc' / 'Projektübersicht.docx'
NAVY='#18384F'; TEAL='#187F83'; GRAY='#586775'; LIGHT='#E8EFF2'; BLACK='#000000'
fontpath=Path('C:/Windows/Fonts')
def f(size,bold=False): return ImageFont.truetype(str(fontpath/('arialbd.ttf' if bold else 'arial.ttf')),size)
def lines(draw,xy,text,width=65,size=30,fill=BLACK,bold=False,gap=10):
    x,y=xy
    for line in textwrap.wrap(text,width=width,break_long_words=False):
        draw.text((x,y),line,font=f(size,bold),fill=fill); y+=size+gap
    return y
def timeline(filename,events):
    im=Image.new('RGB',(1900,180+len(events)*205),'white'); d=ImageDraw.Draw(im)
    x=320; d.line((x,75,x,95+(len(events)-1)*205),fill=TEAL,width=8)
    for i,(date,title,body) in enumerate(events):
        y=65+i*205
        d.ellipse((x-17,y+13,x+17,y+47),fill=TEAL)
        lines(d,(20,y+12),date,width=18,size=34,bold=True,fill=NAVY)
        d.text((380,y),title,font=f(39,True),fill=NAVY)
        lines(d,(380,y+62),body,width=72,size=31,fill=GRAY,gap=9)
    im.save(WORK/filename)
timeline('timeline.png',[
 ('24.09.2026','Erste belegte funktionsfähige Version','v0.2.0 mit Transaktionskern, Fachschicht und Weboberfläche; am selben Abend v0.3.0.'),
 ('10.10.2026','Fachlicher Ablauf und Betriebsfunktionen','v0.4.0 bis v0.5.4: SEPA, Bescheid, Saga, Fachkorrekturen und versionierter Berechnungsstand.'),
 ('11.10.2026','testsweet wird eingeführt','Fach- und Browsermodus, parallele Bearbeitung, SQL-Nachweise und vier Testkategorien.'),
 ('12.10.2026','Breites Review und Nachbesserungen','75 nummerierte Feststellungen; kritische Funde behoben, Monatspläne und Regressionen bis v0.5.8.'),
 ('13.–14.10.','Prüfwerkzeug und Zertifikat werden gehärtet','Vier weitere Primärreviews, echte GLM-Nachprüfung, unabhängige Sollwerte und Quellenbindung.'),
 ('15.10.2026','Weitere Gegenprüfungen und erfolgreicher Abschluss','Fünf Gesamt-Reviews; aktueller Gesamtlauf endet um 22:40 Uhr mit BESTANDEN und 0 Befunden.')])
timeline('testsweet.png',[
 ('11.10.','Abläufe ausführen und Daten prüfen','Simulierte Sachbearbeiter, Browserautomatisierung, Positiv-, Negativ-, Last- und Stressfälle.'),
 ('12.10.','Scheinprüfungen und Seiteneffekte vermeiden','Passwortfreie Berichte, eigene Browserprofile, sichere Wiederaufnahme und globale SQL-Prüfungen.'),
 ('13.10.','Prüfannahmen unabhängig machen','Eindeutige Kombinationen, feste Handwerte und echte GLM-5.3-Prüfung mit vollständigen Dokumenten.'),
 ('14.10.','Gesamtnachweis automatisieren','Neubuild, Pflichtmatrix, GLM-Prüfung und Word-/PDF-Export in einem überwachten Gesamtlauf.'),
 ('15.10.','Quellen und Prüfläufe konsequent isolieren','UUID-Laufkennungen, eingefrorene Runtime, strikter Antwortvertrag, Produktregressionen und Hashkontrollen.')])
im=Image.new('RGB',(1900,650),'white'); d=ImageDraw.Draw(im)
bars=[('24.09. v0.2.0',23),('11.10. Szenario-Regression',136),('12.10. v0.5.5',195),('15.10. Abschlusslauf',395)]
for i,(label,n) in enumerate(bars):
    y=35+i*140; d.text((10,y+12),label,font=f(31),fill=BLACK)
    d.rounded_rectangle((600,y,600+n*2.8,y+65),radius=7,fill=TEAL if i==3 else NAVY)
    d.text((625+n*2.8,y+10),str(n),font=f(37,True),fill=BLACK)
im.save(WORK/'tests.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
sec.header_distance=Cm(.6);sec.footer_distance=Cm(.7)
styles=doc.styles
for st in styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Arial';styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.12
styles['Title'].font.size=Pt(32);styles['Title'].font.bold=True
styles['Title'].paragraph_format.space_after=Pt(8)
styles['Subtitle'].font.size=Pt(14)
styles['Heading 1'].font.size=Pt(21);styles['Heading 1'].paragraph_format.space_after=Pt(12)
styles['Heading 1'].paragraph_format.space_before=Pt(0)
styles['Heading 2'].font.size=Pt(12.5);styles['Heading 2'].paragraph_format.space_before=Pt(11)
styles['Heading 2'].paragraph_format.space_after=Pt(6)
styles['Caption'].font.size=Pt(8.5)
doc.core_properties.title='Projektübersicht'
doc.core_properties.subject='Entwicklung, Reviews und Qualitätssicherung von SDN.Leistung'
doc.core_properties.author='SDN.Leistung Projekt'
doc.core_properties.keywords='Projektgeschichte, Timeline, Reviews, testsweet, Testzertifikat'
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('SDN.Leistung  |  Projektübersicht  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('586775')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def p(text='',bold=False,small=False):
    pp=doc.add_paragraph();r=pp.add_run(text);r.bold=bold
    if small:r.font.size=Pt(9)
    return pp
def h(text,level=2):doc.add_heading(text,level)
def page(title):doc.add_page_break();h(title,1)
def img(file,width=17.3):
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(5)
    pp.add_run().add_picture(str(WORK/file),width=Cm(width))
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for j,a in enumerate(headers): t.rows[0].cells[j].text=a
    for row in rows:
        for j,a in enumerate(row):t.add_row() if False else None
        cells=t.add_row().cells
        for j,a in enumerate(row):cells[j].text=str(a)
    for i,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trPr.append(cs)
        if i==0:
            rep=OxmlElement('w:tblHeader');trPr.append(rep)
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'18384F' if i==0 else ('F1F5F7' if i%2 else 'FFFFFF'));pr.append(sh)
            mar=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                z=OxmlElement('w:'+side);z.set(qn('w:w'),'90');z.set(qn('w:type'),'dxa');mar.append(z)
            pr.append(mar);b=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                z=OxmlElement('w:'+side);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'D9D9D9');b.append(z)
            pr.append(b)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2);pp.paragraph_format.line_spacing=1.04
                if j==0:pp.alignment=WD_ALIGN_PARAGRAPH.LEFT
                for r in pp.runs:r.font.size=Pt(9);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
    doc.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def ref(text):p('Belege: '+text,small=True)

doc.add_paragraph('Projektübersicht','Title')
doc.add_paragraph('SDN.Leistung und Transaktionssteuerung','Subtitle')
p('Entwicklung, Reviewhistorie und Aufbau der Qualitätssicherung',bold=True)
p('Dokumentierter Zeitraum 24. September bis 15. Oktober 2026\nAbschlussstand 15. Oktober 2026 um 22:40 Uhr Europe/Berlin',small=True)
p('Die erste belegte Version war bereits am 24. September lauffähig. In den folgenden drei Wochen wurde sie fachlich erweitert, mehrfach überprüft und um das eigenständige Prüfwerkzeug testsweet ergänzt. Der aktuelle Gesamtlauf besteht alle Pflichtprüfungen mit 0 gemeldeten Fehlern im geprüften Umfang.')
img('timeline.png')
p('10 ausführliche Primärreviews sind für den 12., 13. und 15. Oktober dokumentiert. Frühere Audit- und Gegenprüfungen sind zusätzlich in Versionsänderungen erkennbar. Die aktuelle Ausbaustufe ist als Proof of Concept abgeschlossen; eine behördliche Produktionsfreigabe ist damit nicht verbunden.',small=True)
ref('Q01–Q04, Q05–Q14 und Q23–Q26. Die Grafik ordnet Ereignisse chronologisch; Abstände sind nicht maßstabsgetreu.')

page('Vom ersten Stand zum vollständigen Ablauf')
p('Am 24. September 2026 beginnt die belegbare Projektgeschichte. Der erste Git-Stand enthält schon Version 0.2.0 mit einer funktionsfähigen Transaktionssteuerung und Fachanwendung. Der tatsächliche Beginn der Vorarbeiten ist nicht dokumentiert. Die früheste belegte Version und der Projektbeginn dürfen daher nicht gleichgesetzt werden. [Q01, Q02]')
table(['Datum und Integration','Version','Ergebnis und Bedeutung'],[
 ('24.09. 21:24','0.2.0','Transaktionen, Retry, Tasks, Saga und Oracle; Fachberechnung, Zahlungsworkflow, API und Weboberfläche. 23 Tests im damaligen Stand.'),
 ('24.09. 21:43','0.3.0','Anmeldung mit Rollen, Neuanlage, Kassenlauf und Storno. Aus dem technischen Kern wird ein bedienbarer Arbeitsablauf.'),
 ('10.10. 10:08','0.4.0','SEPA, Vier-Augen-Freigabe, Bescheide und Antrags-Saga; UCP-Pool, Monitoring, Konfiguration, Audit und Handbuch.'),
 ('10.10. 12:47','0.5.0','Fachberechnung erweitert; Pool-, Task- und Saga-Fehler behandelt. Write-Ahead und Recovery sichern mehr Absturzpunkte ab.'),
 ('10.10. 13:15','0.5.1','Externe Audit-Funde zu Titel-XSS, Betragsgrenzen, Jahresfenster und Content Security Policy umgesetzt.'),
 ('10.10. 14:11','0.5.2','Weitere Fachkorrekturen, Monatsprüfung, Bescheidbelehrung und Saga-Kompensation.'),
 ('10.10. 17:49','0.5.3','Bekanntgabe-Text korrigiert; eindeutiger Datenbankindex verhindert mehrere aktive Zahlungen derselben BG und desselben Monats.'),
 ('10.10. 19:41','0.5.4','Rechtsquellen mit Herkunftsnachweisen und eingefrorener Berechnungssnapshot für nachvollziehbare Bescheide.')
],[3.0,1.5,12.8])
h('Was zu diesem Zeitpunkt als fertig gelten konnte')
p('Spätestens am 10. Oktober war der fachliche Ablauf von Antrag über Bescheid bis Auszahlung dokumentiert umgesetzt. Danach lag der Schwerpunkt auf Fehlerbehandlung, Grenzfällen und Nachweisqualität. „Funktionsfähig“ beschreibt den erreichten Funktionsumfang; spätere Reviews zeigen, dass es noch kein Beleg umfassender Fehlerfreiheit war.')
p('Die Uhrzeiten sind Git-Integrationszeiten in MESZ. Sie datieren die Sicherung der Änderungen, nicht den Beginn oder das Ende eines Reviews. Die Versionsfolge ist belegt; eine exakte Dauer jeder Implementierungsphase lässt sich daraus nicht ableiten.',small=True)
ref('Q01: Commits 41f23eb, 0e3c04d, cc6dbf7, 53d9383, 3a4377f, e4d6527, a033638 und cc6da3d; Q03, Q04.')

page('Vom Ablauf zur systematischen Prüfung')
table(['Zeitpunkt','Schritt','Was damit zusätzlich abgesichert wurde'],[
 ('11.10. 09:14','Szenario-Regression','Zehn echte Arbeitsabläufe mehrerer Sachbearbeiter; SQL prüft Beträge, Status, Zuordnungen und Journal direkt im persistierten Endzustand.'),
 ('11.10. 11:13','testsweet Fachmodus','Parallele Mitarbeiter arbeiten einen Auftragstopf ab. GUI, Laufprotokolle sowie PDF- und Markdown-Berichte machen die Prüfung nachvollziehbar.'),
 ('11.10. 14:06','testsweet Webmodus','Selenium führt die Abläufe über Browser und Produktoberfläche aus. Um 14:44 werden drei Wurzelursachen früher Fehlversuche korrigiert.'),
 ('11.10. 22:26','Vier Testkategorien','Positiv, Negativ, Last und Stress. Negativfälle müssen die Ablehnung und den unveränderten Zustand am Abwehrschritt beweisen.'),
 ('12.10. 11:29','v0.5.5 bis v0.5.7','Reviewkorrekturen, Vermögensfortschreibung des Demomodells, Satztabelle, Einsichtsprotokoll und entschärfte Übersichten; gemeinsam integriert.'),
 ('12.10. 11:58–13:25','v0.5.8 und Gegenprüfung','Monatliche Bewilligungspläne, stärkere Sperren, Replay-Bindung und Web-/Audit-/Saga-Härtung. Weitere Korrekturen bei gleicher Version.'),
 ('13.10.','Prüfwerkzeug geprüft','Vier Primärreviews prüfen besonders, ob fehlerhafte Ergebnisse fälschlich als bestanden erscheinen können.'),
 ('14.10.','Gesamtnachweis','Automatisierter Zertifikatslauf verbindet Build, Tests, sieben Simulator-Kombinationen, GLM und Export.'),
 ('15.10. 22:40','Aktuelle Ausbaustufe abgeschlossen','Neuester Gesamtlauf erfolgreich, Zertifikat BESTANDEN. Die nachfolgenden Kapitel erklären Reviews und belastbaren Prüfumfang.')
],[3.0,3.8,10.5])
h('Die wichtigste Veränderung')
p('Die Qualitätssicherung prüft inzwischen auch ihre eigenen Voraussetzungen: korrekte Sollwerte, echte Fehlerrückgaben, unveränderte Quellen und vollständige Nachweise. Ein grünes Ergebnis setzt heute mehr voraus als einen erfolgreich ausgeführten Fachablauf.')
ref('Q01, Q03, Q04, Q15, Q16, Q22–Q26.')

page('Die ersten fünf ausführlichen Reviews')
p('Das Register zählt eigenständige Primärberichte. Teilberichte, Fixberichte und erneute Testläufe werden nicht als zusätzliche Primärreviews gezählt. Am 12. Oktober liegt der erste, bis zum Ende des 13. Oktober liegen fünf solche Berichte vor. Frühere Audits vom 10. Oktober sind in der Versionsgeschichte belegt, aber ohne vollständig rekonstruierbare Einzelbilanz. [Q01, Q03]')
table(['Nr und Datum','Prüfung und KI-Zuordnung','Gemeldete Punkte'],[
 ('1 | 12.10.','Gesamtreview von Anwendung, Kern und testsweet. KI-Modell im Bericht nicht benannt. [Q05]','75 nummerierte Feststellungen: 3 Blocker, 9 kritisch, 24 Major, 24 Minor, 15 Info.'),
 ('2 | 13.10.','testsweet Erstprüfung. KI-Modell nicht belegt. [Q06]','20 bestätigte Befunde: 3 P1, 15 P2, 2 P3.'),
 ('3 | 13.10. abends','Beobachtungsreview von Tests, GLM und Zertifikat. KI-Modell nicht belegt. [Q07]','20 nummerierte Befunde CR101–120: 1 kritisch, 4 hoch, 15 mittel; zusätzlich kompakte niedrige Beobachtungen.'),
 ('4 | 13.10.','Gesamtprüfung der Nachweiskette. KI-Modell nicht belegt. [Q08]','16 Befunde R01–R16: 13 P1, 3 P2; Überschneidungen mit dem Abendbericht.'),
 ('5 | 13.10. zweite Prüfung','Statische Fixverifikation durch ZCode mit GLM-5.3, ausdrücklich im Bericht benannt. [Q09]','10 neue nummerierte Punkte CR201–210: 3 mittel, 7 niedrig.')
],[2.4,7.1,7.8])
h('Bilanz bis zum fünften Review')
p('Der erste Review enthält 60 problemorientierte Punkte und 15 Info-Feststellungen. Die vier Berichte vom 13. Oktober enthalten zusammen 66 nummerierte Hauptbefundnennungen sowie 30 zusätzliche niedrige Beobachtungen im Abendbericht. Diese Mengen sind keine Zahl verschiedener Fehler: Die Berichte überschneiden sich und bewerten zum Teil dieselben Nachweisprobleme.')
p('Im Bericht vom 12. Oktober nennt die Summenzeile 23 Minor-Punkte, die ID-Liste CR037–060 enthält jedoch 24. Im zweiten Bericht vom 13. Oktober verweist die Einleitung bis CR213, tatsächlich ausgeführt sind CR201–210. Das Register folgt den konkreten ID-Listen. Info-Punkte enthalten auch positive Bestätigungen und zählen nicht automatisch als Fehler.',small=True)
ref('Q05–Q09. P1/P2/P3 und Blocker/Kritisch/Major/Minor sind unterschiedliche Bewertungsmaßstäbe.')

page('Fünf weitere Gesamtprüfungen am 15 Oktober')
p('Mit den fünf Primärberichten vom 15. Oktober wächst das Register auf zehn ausführliche Reviews. Die KI-Namen sind nur dort angegeben, wo zugehörige Gegenprüfungsberichte die Herkunft ausdrücklich zuordnen. Die dort verwendeten Kurzkennungen identifizieren den Bericht, nicht die Modellversion. [Q10–Q14, Q18–Q21]')
table(['Nr und Bericht','KI soweit belegt','Priorisierte Befundnennungen'],[
 ('6 | 3c909c2 [Q10]','Nicht belegt','10: 2 kritisch und 8 Major. Storno-Rechte, Fachrechnung, Recovery, Snapshot und Prüfvertrag.'),
 ('7 | k4p9 [Q11]','Grok; genaue Modellversion nicht belegt','16: 3 kritisch, 6 Major, 7 Minor. Abbruchstatus, Saga-Namen, Kassenlauf, Vorschau und weitere Punkte.'),
 ('8 | m9x2 [Q12]','Antigravity als Werkzeug genannt; zugrunde liegendes Modell nicht belegt','2: 1 Major, 1 Minor. Testisolation und Nachweisdokumentation; vorgeschlagene Fachreparatur wird teilweise widerlegt.'),
 ('9 | r7w2 [Q13]','DeepSeek; genaue Modellversion nicht belegt','8: 3 kritisch, 2 Major, 3 Minor. Commitquittung, Lastaussage, Cleanup und Eingabevertrag; auch bekannte Modellgrenzen.'),
 ('10 | zc1 [Q14]','GLM-5.3 laut zugehöriger Gegenprüfung','5 priorisierte Punkte: 1 Major, 4 Minor; zusätzlich 9 Info-Beobachtungen in Gruppe ZC-6.')
],[4.0,4.5,8.8])
h('Bilanz nach zehn Primärreviews')
p('Am 15. Oktober wurden in den fünf Gesamtberichten 41 priorisierte Befundnennungen und zusätzlich neun Info-Beobachtungen erfasst. Ein Teil betrifft bereits bekannte Grenzen oder wird durch Gegenprüfung präzisiert beziehungsweise widerlegt. Eine Addition mit den älteren Berichten ergäbe daher keine belastbare Zahl verschiedener Softwarefehler.')
table(['Stichtag','Primärreviews kumuliert','Befundbilanz der jeweiligen Welle'],[
 ('12.10.','1','75 Feststellungen, davon 60 problemorientiert und 15 Info'),
 ('13.10.','5','66 nummerierte Hauptbefundnennungen plus 30 niedrige Beobachtungen'),
 ('15.10.','10','41 priorisierte Befundnennungen plus 9 Info in fünf weiteren Berichten')
],[3.0,4.3,10.0])
p('Die Reihenfolge innerhalb desselben Tages ist nicht vollständig als Start-/Endzeitfolge belegt. Die Nummerierung dient dem Register. Die abnehmende Zahl pro Welle ist wegen unterschiedlicher Umfänge und Doppelmeldungen keine statistische Fehlerquote.',small=True)

page('Wie aus Reviewfunden dauerhafte Korrekturen wurden')
h('12 Oktober Datenschutz und Fachwege')
p('Die drei Blocker des ersten Gesamtreviews betrafen Sozialdaten im Audit-Log, Passwörter in testsweet-Berichten und abweichende Berechnungen zwischen Saga und manuellem Fachweg. Die Korrekturen kürzten Logs auf zulässige Metadaten, entfernten Passwörter aus Berichten und banden die Saga-Berechnung an denselben Datenbestand und dieselben Fachparameter. [Q05, Q15]')
p('Alle drei Blocker und neun kritischen Funde wurden als behoben dokumentiert. Die Hauptsuite wuchs beim ersten Fixstand von 136 auf 195 Tests. Spätere Ergänzungen führten bis zur gemeinsamen Integration von v0.5.5–v0.5.7 auf 213 Tests. Unterschiedliche historische Teststände werden nicht miteinander vermischt. [Q01, Q15]')
h('12 Oktober Monatspläne und Konkurrenzfälle')
p('Eine zusätzliche Nachprüfung erfasste elf Punkte N01–N11 und dokumentierte deren Korrektur. Dazu gehören monatsgenaue Bescheide, Ende des Schwangerschaftsmehrbedarfs, bekannte Satzwechsel, Sperren gegen gleichzeitige Historisierung und Zahlungs-Replays mit Bindung an BG, Monat und Bankdaten. Die Einsichtsprotokollierung wurde transaktional dedupliziert. [Q16]')
h('15 Oktober zehn konsolidierte Gesamtfunde')
p('Die zehn GR-Funde wurden im Fixbericht einzeln mit Korrektur und Regression verknüpft. Besonders folgenreich waren die serverseitige Adminpflicht beim Storno, der günstigere Alleinerziehendenmehrbedarf, der entfallende zentrale Warmwasserabzug und ein vollständig eingefrorenes Bescheiddokument. Hinzu kamen Login-Schutz, sitzungsgebundene Drucktoken, strikte Java-/Node-Antwortparität und vendorcodegebundene Retries. [Q10, Q17]')
table(['Konkreter Fund','Korrektur und dauerhafte Gegenprobe'],[
 ('Storno umgeht Rechteprüfung','Rollenprüfung unmittelbar vor dem Serviceaufruf. HTTP→Service→Oracle-Gegenproben prüfen alternative ID-Schreibweisen.'),
 ('Falsche gemeinsame Sollannahme','Feste, vorab handgerechnete Fälle für vier bis sechs Kinder sowie nicht null gesetzten Warmwasser-Altparameter.'),
 ('Storno-Scheinprüfung','Echte SELECT-Abfragen lesen Vermögen und Restbudget unmittelbar nach Storno vor der nächsten Mutation; Fehler stoppen den Auftrag.'),
 ('Browser nutzt veraltete Listenzeile','Nur lesendes Öffnen erhält einen frischen Locator. Regressionen verhindern Wiederholung bestätigter Schreibschritte.')
],[5.2,12.1])
ref('Q15–Q17. Für das ursprüngliche Gesamtregister gibt es keine vollständig entdoppelte Abschlussliste aller 60 problemorientierten Punkte.')

page('Gegenprüfungen bestätigten auch Grenzen')
p('Nicht jede gemeldete Schwäche war ein neuer Implementierungsfehler, und nicht jede vorgeschlagene Reparatur war fachlich richtig. Die Gegenprüfungen hielten diese Unterschiede fest und verhinderten, dass unbelegte Änderungen als Fehlerbehebung erscheinen.')
table(['Ausgangsreview','Ergebnis der Gegenprüfung','Konkrete Wirkung'],[
 ('Grok | 16 Punkte','12 Korrekturen, 3 Reparaturforderungen nicht bestätigt, 1 offene Zuständigkeitsgrenze. [Q18]','Unbekannter Task-Ausgang statt Rollbackzusage; persistierte Saga-Namen; Kassenupdates in 1.000er-Blöcken; Monatsvorschau, Änderungsbescheid und Eingabevalidierung.'),
 ('Antigravity | 2 priorisierte Punkte','Isolationslücken und Berichtslücken korrigiert; behauptete Simulator-Ursache und vorgeschlagene Bescheidpflicht nicht als bewiesen übernommen. [Q19]','Eigene UUIDs und Sequenzen, getrenntes Cleanup, eindeutige Journalerklärung; laufgebundene Runtime schützt vor parallelen Buildänderungen.'),
 ('GLM-5.3 | 5 priorisierte Punkte','Betriebsbefund und kleinere Punkte adressiert; Testlücke als Testlücke eingeordnet. Zusätzlich Integritätsrisiko historischer Journale entdeckt. [Q20]','Offline-Werkzeug und Wartungsanleitung; namenlose historische PENDINGs bleiben konservativ ungeklärt. Monatsbezug der Liste, Nullmapping und Druckrouting korrigiert.'),
 ('DeepSeek | 8 priorisierte Punkte','Teils bestätigt und korrigiert, teils bekannte Grenzen oder nicht tragfähige Reparaturforderungen. [Q21]','Commitquittung als unbekannter Ausgang, auch bei fatalem Treiberfehler; UUID-Cleanup, Subcent-/Schlüsselvalidierung, Perzentile und verpflichtende Node-Prüfer.')
],[3.3,6.3,7.7])
h('Drei verbleibende Modellgrenzen')
p('Objekt- und Teamzuständigkeiten sind weiterhin nicht implementiert. Eine beleggebundene Vermögenshistorie fehlt; ein rein rechnerischer Verbrauch darf nicht als allgemeine rechtliche Lösung ausgegeben werden. Ein strenger Betriebsmodus mit durchgehender Bescheidpflicht und vollständigem Verwaltungsverfahren bleibt eine gesonderte Erweiterung. Diese Grenzen werden durch ein grünes Testzertifikat nicht aufgehoben. [Q18–Q21]')
p('Altjournale benötigen besondere Vorsicht: Ein namenloses PENDING aus einer frühen Engine beweist keine Wirkungslosigkeit. Die heutige Schutzlogik behält ungeklärte Daten und verlangt eine kontrollierte Klärung. Das verhindert eine scheinbar erfolgreiche Kompensation auf ungesicherter Grundlage. [Q20]')

page('Die Entwicklung von testsweet')
p('testsweet wurde am 11. Oktober als eigenständiger Mitarbeiter-Simulator eingeführt. Bis zum Abschluss am 15. Oktober entwickelte sich daraus eine mehrstufige Nachweiskette, die Fachabläufe, Weboberfläche, Datenbankzustände und die Verlässlichkeit der Prüfer verbindet. [Q01, Q22]')
img('testsweet.png')
h('Breitere Fälle statt Wiederholung derselben Konstellation')
p('Der Kombinationsgenerator ergänzt kuratierte Fälle um unterschiedliche Haushalte, Merkmale, Einkommen, Vermögen, Wege und Verstoßkombinationen. Eine Signaturmenge verhindert Wiederholungen. Feste Fachgrenzfälle sichern besonders wichtige Konstellationen unabhängig von Zufallsseeds ab und werden für die GLM-Stichprobe bevorzugt. [Q03, Q22]')
ref('Q01, Q03, Q06–Q09 und Q22. Die Grafik zeigt Ausbaustufen; einzelne Funktionen wurden anschließend weiter korrigiert.')

page('Die Prüfkette wurde selbst zum Prüfgegenstand')
h('13 Oktober echte Modellprüfung und strengere Urteile')
p('Alle 20 TS-Befunde des ersten testsweet-Reviews wurden laut Nachtrag korrigiert und gegengeprüft. Eine Python-Nachrechnung wurde nicht länger als LLM-Prüfung bezeichnet. Die echte GLM-API erhielt vollständige benannte Prüfdokumente und lieferte gespeicherte Antworten. Transportnachweis und fachliches Urteil wurden getrennt. [Q06, Q22]')
p('Die nachfolgenden Reviews CR101–120 und R01–R16 wurden gemeinsam bearbeitet. Separate BigDecimal-/BigInt-Referenzen, vollständige Zustandsvergleiche für Negativfälle und eine Fehleraggregation mit Vorrang von FAIL schlossen konkrete Wege zu falschem Grün. Fehlende, übersprungene, abgebrochene oder veraltete Pflichtprüfungen dürfen kein positives Zertifikat erzeugen. [Q07, Q08, Q27]')
table(['Ausbauschritt','Neu gewonnene Absicherung'],[
 ('Kombinationsgenerator','Garantiert unterschiedliche Signaturen statt bloßer Wiederholung kuratierter Fälle; kontrollierte Haushalts-, Merkmals- und Verstoßkombinationen.'),
 ('Feste Fachgrenzfälle TF11–TF19','Vorab vorgegebene Sollwerte und bevorzugte Auswahl für GLM. Spätere Reviews verschärften Kinder- und Warmwasserfälle erneut.'),
 ('Prüfdokumente und Snapshots','Prüfauftrag, Rechtsgrundlagen und vollständige Fallgruppen mit Manifest, Quell-/Request-/Response-Snapshots und SHA-256.'),
 ('Antwortvertrag und Parität','Striktes JSON, vollständige Zahlen und eindeutige Fallzuordnung. Gemeinsamer Java-/Node-Rohbyteskorpus erkennt Parserunterschiede.'),
 ('Build- und Quellenbindung','JUnit, Simulator, GLM und Zertifikat müssen zum gleichen erfassten Quellenstand gehören. Ein Nachweis wird bei Änderungen veraltet.'),
 ('15.10. zusätzlicher Produktschutz','28 zusätzliche Java-Ausführungen und 3 JavaScript-Tests stärken HTTP-, Oracle-, Task-, Saga-, Kassen- und UI-Gegenproben.')
],[5.1,12.2])
h('14 Oktober erster erfolgreicher Gesamtlauf')
p('Der damalige automatisierte Gesamtlauf von 01:44 bis 01:57 Uhr bestand 306 Haupttests, 155 testsweet-Tests, 1.312 Simulatoraufträge und 18 GLM-Fälle in sechs Gruppen. Er verband erstmals die vollständige Ein-Klick-Prüfkette mit dem Zertifikatsexport. Diese Zahlen gelten für den damaligen Quellenstand. [Q28]')
h('15 Oktober größerer verbindlicher Abschlussumfang')
p('Standardläufe wurden auf 100 positive und 100 negative Fälle angehoben; die GLM-Standardauswahl umfasst genau 100/100 ohne einseitige Auffüllung. Web-Zertifikatsläufe stiegen von 60 auf 100 Fälle, während die Fachbreite bei 500/500 blieb. Daraus ergibt sich die aktuelle Pflichtmatrix von 1.392 Simulatoraufträgen. [Q22, Q24]')

page('Wie die Fehlersicherheit schrittweise zunahm')
h('1 Datenbank und Transaktionssteuerung')
p('Oracle übernimmt atomare Änderungen, Sperren und Recovery. Der Steuerungskern ergänzt kurze Transaktionen, kontrollierte Wiederholung ausschließlich bei Sperrkonflikten, Zeitlimits und Saga-Journale. Constraints und der Monatsindex bilden zusätzliche Barrieren gegen inkonsistente Zahlungen. Ein unklarer Commit-Ausgang wird als unklar gemeldet und nicht blind wiederholt. [Q03, Q18, Q21]')
h('2 Regressionen nach jedem bestätigten Fund')
p('Bestätigte Probleme wurden durch gezielte Tests wieder erreichbar gemacht: verlorene Commitquittung, alternative Storno-IDs, Definitionswechsel einer Saga, konkurrierende Änderungen und echter Oracle-Endzustand. Die Tests prüfen die verletzte Eigenschaft; ein erfolgreicher Build allein würde sie nicht belegen.')
img('tests.png')
p('Entwicklung der registrierten Haupttests an ausgewählten historischen Ständen. Die Fälle und der Umfang änderten sich; Testzahl ist ein Ausbauindikator und keine gemessene Fehlerfreiheit. [Q01, Q05, Q15, Q24]',small=True)
h('3 Unabhängige Sollwerte und negative Gegenproben')
p('Handwerte und separate Referenzrechnungen ergänzen den Produktcode. Negativfälle müssen einen konkreten Fachfehler auslösen und den Zustand am Abwehrschritt unverändert lassen. Der Review vom 15. Oktober zeigte dennoch gemeinsame falsche Fachannahmen in mehreren Referenzen. Deshalb wurden gerade diese Grenzfälle anschließend mit extern begründeten Sollwerten ergänzt. [Q06, Q10, Q17]')
h('4 Prüfung der Nachweise selbst')
p('Quellen- und Buildhashes, eindeutige Testnamen, vollständige Berichte, Laufisolation und ein strikter Modellantwortvertrag schützen das Gesamturteil. Fehlerhafte, abgebrochene oder veraltete Pflichtnachweise verhindern ein positives Zertifikat. Das reduziert konkret bekannte Fehlermöglichkeiten; eine verbleibende Fehlerwahrscheinlichkeit wurde nicht statistisch bestimmt. [Q22–Q26]')

page('Fehlversuche blieben sichtbar und wurden aufgearbeitet')
p('Die Historie besteht nicht aus ausschließlich grünen Läufen. Mehrere Fehlschläge führten zu zusätzlichen Regressionen oder strengeren Nachweisen. Die ursprünglichen Belege blieben erhalten. Das ist entscheidend, um Verbesserungen nachvollziehen zu können.')
table(['Datum','Beobachtetes Ergebnis','Konsequenz'],[
 ('13.10.','Ein Negativlauf mit anderem Seed erreichte 499/500; ein Weblauf mit einem Mitarbeiter 0/1 trotz fertiger Phasenmeldung. [Q06]','Testdaten, Rollenwechsel und Status-/Fehlerauswertung verbessert. Erwartete Exceptions müssen tatsächlich auftreten.'),
 ('15.10. erste GR-Fixprüfung','GLM zunächst 17 PASS und 1 UNVERIFIED: tatsächlich verwendete SEPA-Rohwerte fehlten im Prüfmaterial. [Q17]','Rohwerte aus derselben Variablenquelle wie die Aktion in Berichte aufgenommen; neue Regressionen vergleichen Übergabe und Dokumentation.'),
 ('15.10. Browser-Fixprüfung','59/60 erfolgreiche Aufträge; ersetzte DOM-Zeile beim lesenden Öffnen. [Q17]','Gezielte Wiederholung nur des lesenden Locators; Schreibaktionen bleiben ohne automatische Wiederholung.'),
 ('15.10. 13:44','testsweet-Phase: 288 Tests, 5 SQL-Ausnahmen wegen fehlender Demo-Spalte SAGA_ID im bestehenden Schema. [Q23]','Fixtures vom neuen DDL entkoppelt. Wiederholung 14:03–14:05: 291 Tests, 0 Fehler, 0 Skips. Gesamtlauf damals noch nicht wiederholt.'),
 ('15.10. nachmittags','Ein Gesamtlauf blieb trotz 1.392 bestätigten Simulatoraufträgen unverifiziert: Produkt-JAR wurde parallel ersetzt. [Q19, Q20]','Gehasht kopierte Runtime und fortlaufende Kontrolle der Originalbelege. Die Quellenbindung wurde nicht gelockert.'),
 ('15.10. 21:29–22:40','Neuester Gesamtlauf: SUCCESS, alle 13 Phasen erfolgreich, 0 Findings. [Q24–Q26]','Erneuter vollständiger Nachweis einschließlich GLM und veröffentlichtem Word-/PDF-Zertifikat.')
],[2.9,7.0,7.4])
h('Warum ältere Fehlschläge den aktuellen Abschluss nicht entwerten')
p('Die Korrekturen wurden anschließend am aktuellen Quellenstand erneut geprüft. Frühere Fehlversuche erklären die zusätzliche Absicherung; sie wurden nicht zu Erfolgen umgedeutet. Umgekehrt ersetzen ältere grüne Ergebnisse keinen aktuellen quellengebundenen Gesamtlauf.')

page('Abschlussnachweis am 15 Oktober 2026')
p('Der jüngste Zertifikatsgesamtlauf startete um 21:29:01 Uhr und endete um 22:40:32 Uhr Europe/Berlin. Alle 13 Phasen wurden erfolgreich abgeschlossen. Der veröffentlichte Zustand lautet READY mit Urteil BESTANDEN und 0 Findings. [Q24–Q26]')
table(['Prüfebene','Ergebnis','Aussage'],[
 ('Hauptprojekt JUnit','395 / 395','0 Failures, 0 Errors, 0 Skips'),
 ('testsweet JUnit','329 / 329','0 Failures, 0 Errors, 0 Skips'),
 ('Fachmodus Positiv','500 / 500','Persistierte Sollzustände bestätigt'),
 ('Fachmodus Negativ','500 / 500','Erwartete Fachabwehr und Zustände bestätigt'),
 ('Fachmodus Last','120 / 120','Begrenzte funktionale Parallelprüfung'),
 ('Fachmodus Stress','42 / 42','21 Fälle in zwei Wellen'),
 ('Webmodus Positiv','100 / 100','Browserabläufe mit Endzustandsprüfung'),
 ('Webmodus Negativ','100 / 100','Abwehrfälle über die Oberfläche'),
 ('Webmodus Last','30 / 30','Begrenzte Browser-Parallelprüfung'),
 ('Simulator gesamt','1.392 / 1.392','Sieben Pflichtkombinationen, 0 Fehlaufträge'),
 ('GLM-5.3','200 / 200 PASS','100 positive und 100 negative Fälle, 67 Gruppen, 0 Findings')
],[5.1,3.4,8.8])
p('Die 200 GLM-Fälle sind eine Auswahl aus Simulatorberichten und zählen nicht als weitere unabhängige Simulatoraufträge. Ebenso werden Java-, Node- und Python-Prüfzahlen nicht zu einer gemeinsamen Qualitätsquote addiert.',small=True)
h('Auch die formalen Modellantworten wurden geprüft')
p('59 der 67 GLM-Gruppen waren beim ersten Versuch gültig, sieben nach einer Wiederholung und eine nach zwei Wiederholungen. Insgesamt gab es neun erneute Sendungen und keine endgültig gescheiterte Gruppe. Formal ungültige Antworten wurden mit protokolliertem Korrekturhinweis erneut angefordert. Die Rohantworten blieben nachvollziehbar. [Q25]')
p('Das Ergebnis lautet damit präzise: 0 beobachtete Fehler beziehungsweise Befunde im dokumentierten Prüfumfang des aktuellen Gesamtlaufs. Es ist kein mathematischer Nachweis, dass im gesamten System keine weiteren Fehler existieren.')

page('Was der Abschluss belegt')
h('Aktuelle Ausbaustufe')
p('Die technische und fachliche Demonstrationsanwendung ist implementiert und lauffähig. testsweet deckt jetzt Fach- und Webwege, positive und negative Abläufe, begrenzte Parallelprüfungen, externe Modellnachrechnung und den automatisierten Zertifikatsexport ab. Die aktuelle Ausbaustufe endet am 15. Oktober mit einem vollständigen erfolgreichen Nachweis. [Q04, Q22–Q26]')
h('Bindung an den tatsächlich geprüften Stand')
p('Der Abschlussnachweis gehört zum Arbeitsbaum auf Basiscommit 37abfb6 und enthält uncommittierte Änderungen. Der Commit allein genügt deshalb nicht zur Reproduktion. Das Quelleninventar erfasst 288 Dateien. Eine erneute SHA-256-Kontrolle bei Erstellung dieser Übersicht ergab keine Abweichung; auch die fünf im Publikationsmanifest ausgewiesenen Dateien stimmen mit ihren gespeicherten Prüfsummen überein. [Q24–Q26]')
table(['Kennung','Wert'],[
 ('Gesamtlauf','3a0bbe85-bd68-4739-8674-fd92b3329223'),
 ('Zertifikat','b805c5b6-20b0-4215-83a5-26900a2f31e4'),
 ('Quellenfingerabdruck','3b3836e2d868ea8ac81b7bdd7b77ad43c\ne75a8d6e2c4b8e961574f5dd491741b'),
 ('Evidenzfingerabdruck','5ee667ff6075790fd1b2b99f3db332207\n393b53524760181e42fd7325dae57ca')
],[4.8,12.5])
h('Praktische Grenzen der Aussage')
p('Der aktuelle Nachweis wurde mit Temurin 25.0.4.1 und Compilerziel Java 21 sowie Oracle Free 23.26.3 erzeugt. Er belegt keinen erneuten Gesamtlauf auf einer Java-21-Laufzeit oder Oracle 19c Enterprise Edition. Historische Prüfungen anderer Umgebungen bleiben davon getrennt. [Q24, Q26]')
p('LAST und STRESS belegen funktionale Parallelabläufe im ausgeführten Umfang. Sie sind kein Dauerlast-, Sättigungs- oder Kapazitätsnachweis. GLM liefert eine zusätzliche Fallprüfung und keine amtliche fachliche Freigabe. Die bekannten Grenzen bei Zuständigkeiten, Vermögensnachweisen und Verwaltungsverfahren bestehen fort. [Q18–Q22]')
h('Ergebnis der Entwicklung')
p('Die nachvollziehbare Verbesserung liegt in den geschlossenen konkreten Fehlerpfaden und ihren Regressionen, den breiteren Szenarien sowie einer strengeren Kontrolle der Nachweise. Eine lückenlose Gesamtzahl aller jemals gefundenen und behobenen Einzeldefekte ist nicht rekonstruierbar. Dokumentiert sind zehn ausführliche Primärreviews, weitere Gegenprüfungen und der erfolgreiche Abschlusslauf.')

sources=[
 ('Q01','Git-Historie','40 Commits vom 24.09. bis 15.10.2026; git log --reverse. Frühester Stand 41f23eb, zuletzt 37abfb6.'),
 ('Q02','Projektbeschreibung','doc/projektbeschreibung.md; Ausgangsbeschreibung vom 24.09.2026.'),
 ('Q03','Roadmap','doc/roadmap.md; erreichte Etappen und ältere Reviewkorrekturen.'),
 ('Q04','Projektstatus','README.md und AGENTS.md; Proof of Concept und dokumentierte Modellgrenzen.'),
 ('Q05','Gesamtreview 12.10.','doc/codereview/codereview-bericht-2026-10-12.md'),
 ('Q06','testsweet Erstprüfung','doc/codereview/testsweet-codereview-2026-10-13.md'),
 ('Q07','testsweet Abendprüfung','doc/codereview/testsweet-glm-zertifikat-codereview-2026-10-13.md'),
 ('Q08','testsweet Gesamtprüfung','doc/codereview/testsweet-gesamtpruefung-2026-10-13.md'),
 ('Q09','testsweet Fixverifikation','doc/codereview/testsweet-codereview-2026-10-13b-fixverifikation.md'),
 ('Q10','Gesamtprüfung 3c909c2','doc/codereview/grundsicherung-gesamt-review-2026-10-15-3c909c2.md'),
 ('Q11','Grok k4p9','doc/codereview/grundsicherung-gesamt-review-2026-10-15-k4p9.md'),
 ('Q12','Antigravity m9x2','doc/codereview/grundsicherung-gesamt-review-2026-10-15-m9x2.md'),
 ('Q13','DeepSeek r7w2','doc/codereview/grundsicherung-gesamt-review-2026-10-15-r7w2.md'),
 ('Q14','GLM zc1','doc/codereview/grundsicherung-gesamt-review-2026-10-15-zc1.md'),
 ('Q15','Fixbericht 12.10.','doc/codereview/codereview-fixbericht-2026-10-12.md'),
 ('Q16','Nachbesserung 12.10.','doc/codereview/codereview-nachbesserungsbericht-2026-10-12.md'),
 ('Q17','GR-Korrekturen 15.10.','doc/codereview/grundsicherung-fixbericht-2026-10-15.md'),
 ('Q18','Grok Gegenprüfung','doc/codereview/grundsicherung-grok-gegenpruefung-2026-10-15.md; ergänzende Kern-, Fach- und Web-Gegenprüfungen.'),
 ('Q19','Antigravity Gegenprüfung','doc/codereview/antigravity-gegenpruefung-2026-10-15.md'),
 ('Q20','GLM Gegenprüfung','doc/codereview/glm53-gegenpruefung-2026-10-15.md'),
 ('Q21','DeepSeek Gegenprüfung','doc/codereview/deepseek-gegenpruefung-2026-10-15.md; ergänzend deepseek-kern-gegenpruefung-2026-10-15.md.'),
 ('Q22','testsweet und Abdeckung','testsweet/README.md; doc/codereview/testsweet-abdeckung-2026-10-15.md; doc/testzertifikat-werkzeug/README.md.'),
 ('Q23','Saga-Testfixture','doc/codereview/zertifikat-saga-testkorrektur-2026-10-15.md und gleichnamiger -beleg.json.'),
 ('Q24','Aktueller Prüfstand','doc/testzertifikat-werkzeug/stand.json; Schichten, Läufe, Quelleninventar und Fingerabdrücke.'),
 ('Q25','Aktueller Gesamtlauf','testsweet/logs/certificate-suite-3a0bbe85-bd68-4739-8674-fd92b3329223/status.json; darin referenzierte Lauf- und GLM-Belege.'),
 ('Q26','Veröffentlichtes Zertifikat','doc/testzertifikat-werkzeug/current-certificate.json; doc/SDN-Leistung-LMG-Testzertifikat.docx und .pdf.'),
 ('Q27','Gemeinsame testsweet-Fixes','doc/codereview/testsweet-review-fixes-2026-10-13.md; testsweet-review-verifikation-2026-10-13.json.'),
 ('Q28','Erster Zertifikatsgesamtlauf','doc/codereview/testsweet-zertifikatslauf-verifikation-2026-10-14.json; Suite 43ae02a2-c623-48d6-9f53-c1ea8ede1ef0.')]
# Resolve the standalone R01 report if its file name differs from its folder.
r_candidates=list((ROOT/'doc/codereview').glob('*.md'))
for pp in r_candidates:
    txt=pp.read_text(encoding='utf-8-sig')
    if 'R01' in txt and 'R16' in txt and 'gesamtpruefung' in pp.name:
        sources[7]=('Q08','testsweet Gesamtprüfung',pp.relative_to(ROOT).as_posix())
page('Quellen und Zählregeln')
p('Die Quellenkennungen verbinden Aussagen mit den Projektunterlagen. Pfade beziehen sich auf das Repository LeistungsmanagementGrundsicherung. Der vorliegende Bericht ist eine historische Projektübersicht; maßgeblich für den aktuellen Prüfstatus bleiben das veröffentlichte Zertifikat und seine maschinellen Nachweise.')
table(['Quelle','Nachweis'],[(a,b+'\n'+c) for a,b,c in sources[:14]],[1.5,15.8])
p('Reviewzahlen wurden aus den nummerierten Befundlisten übernommen und bei Widersprüchen zur Summenzeile erläutert. Prüfernamen wurden nicht aus Git-Autoren, Dateinamen-Kürzeln oder Schreibstil abgeleitet. Modellnachrechnung mit GLM ist von einem durch GLM verfassten Codereview zu unterscheiden.',small=True)
page('Weitere Belege und Leseschlüssel')
table(['Quelle','Nachweis'],[(a,b+'\n'+c) for a,b,c in sources[14:]],[1.5,15.8])
p('Datumsangaben beziehen sich auf Europe/Berlin. Review-Uhrzeiten werden nur genannt, wenn sie in Laufbelegen stehen. Historische Testzahlen beschreiben den damaligen Stand. Eine vollständige Entdoppelung aller Befunde über alle Berichte ist nicht belegt; deshalb wird keine fiktive Gesamtfehlerzahl ausgewiesen.',small=True)

DOC.parent.mkdir(exist_ok=True);doc.save(DOC)
manifest={}
for pp in (ROOT/'doc/codereview').glob('*.md'):
    manifest[pp.relative_to(ROOT).as_posix()]=hashlib.sha256(pp.read_bytes()).hexdigest()
for path in ['README.md','AGENTS.md','doc/roadmap.md','doc/projektbeschreibung.md','doc/testzertifikat-werkzeug/stand.json','doc/testzertifikat-werkzeug/current-certificate.json']:
    manifest[path]=hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
(WORK/'source-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
print(str(DOC))
